"""Optional local XLSX-to-CSV preprocessor for the financial Skills.

This module is deliberately outside the Skill ZIPs.  It reuses the document
engine's path, container, snapshot, and spreadsheet bounds instead of adding an
Excel runtime to each Skill.  WorkBuddy/platform runtime availability is not
inferred by this local helper.
"""
from __future__ import annotations

import argparse
import csv
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable, Sequence

try:
    from wangbin_document_engine import SecurityPolicy, SpreadsheetService
    from wangbin_document_engine.errors import EngineError, ExportFailure
    from wangbin_document_engine.metadata import publish_new_file, safe_output_path
except ModuleNotFoundError:  # Allow direct execution from this repository only.
    _engine_src = Path(__file__).resolve().parents[1] / "services" / "document-engine" / "src"
    if str(_engine_src) not in sys.path:
        sys.path.insert(0, str(_engine_src))
    from wangbin_document_engine import SecurityPolicy, SpreadsheetService
    from wangbin_document_engine.errors import EngineError, ExportFailure
    from wangbin_document_engine.metadata import publish_new_file, safe_output_path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

MAX_CSV_BYTES = 1 * 1024 * 1024


def _reject_macros(path: Path) -> None:
    """Reject macro-bearing content even when it has been renamed to .xlsx."""
    try:
        with zipfile.ZipFile(path) as archive:
            names = {item.filename.replace("\\", "/").casefold() for item in archive.infolist()}
            if any(
                name.endswith("/vbaproject.bin")
                or name.endswith("/vbadata.xml")
                or "/macrosheets/" in f"/{name}"
                for name in names
            ):
                raise ExportFailure("MACRO_INPUT", "不支持包含宏的 XLSX 文件")
            content_types = "[content_types].xml"
            if content_types in names:
                payload = archive.read(next(item.filename for item in archive.infolist() if item.filename.casefold() == content_types))
                lowered = payload.lower()
                if b"vbaproject" in lowered or b"macroenabled" in lowered:
                    raise ExportFailure("MACRO_INPUT", "不支持包含宏的 XLSX 文件")
    except zipfile.BadZipFile as exc:
        raise ExportFailure("CORRUPT_CONTAINER", "XLSX 文件容器损坏") from exc


def _formula_cell(cell: Cell) -> bool:
    value = cell.value
    if cell.data_type == "f":
        return True
    # Reject formula-injection-shaped text too; numeric negative values are not
    # strings and therefore remain valid financial values.
    return isinstance(value, str) and value.startswith(("=", "+", "-", "@"))


def _scan_rows(worksheet: object, *, require_value: bool = True) -> tuple[list[tuple[object, ...]], bool]:
    """Read the selected bounded worksheet and reject formulas before export."""
    rows: list[tuple[object, ...]] = []
    has_value = False
    for row in worksheet.iter_rows():
        values: list[object] = []
        for cell in row:
            if _formula_cell(cell):
                raise ExportFailure("FORMULA_INPUT", "不支持包含公式或公式样式文本的输入")
            if cell.value is not None:
                has_value = True
            values.append(cell.value)
        rows.append(tuple(values))
    if require_value and not has_value:
        raise ExportFailure("EMPTY_SHEET", "选定工作表为空")
    return rows, has_value


def _scan_workbook(workbook: object, sheet: str) -> list[tuple[object, ...]]:
    selected_rows: list[tuple[object, ...]] | None = None
    for worksheet in workbook.worksheets:
        rows, _ = _scan_rows(worksheet, require_value=worksheet.title == sheet)
        if worksheet.title == sheet:
            selected_rows = rows
    if selected_rows is None:
        raise ExportFailure("SHEET_NOT_FOUND", "指定工作表不存在")
    return selected_rows


def _selected_sheet_names(summary: object) -> set[str]:
    return {str(item["name"]) for item in summary.worksheets}


def _write_csv(rows: Iterable[Sequence[object]], temporary: Path) -> None:
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for row in rows:
            writer.writerow(row)
            handle.flush()
            if temporary.stat().st_size > MAX_CSV_BYTES:
                raise ExportFailure("EXPORT_TOO_LARGE", "CSV 导出文件超过 1 MiB 限制")
    if not temporary.exists() or temporary.stat().st_size == 0:
        raise ExportFailure("INVALID_EXPORT", "CSV 导出文件为空")
    if temporary.stat().st_size > MAX_CSV_BYTES:
        raise ExportFailure("EXPORT_TOO_LARGE", "CSV 导出文件超过 1 MiB 限制")
    with temporary.open("r", encoding="utf-8", newline="") as handle:
        next(csv.reader(handle), None)


def convert_xlsx_to_csv(
    input_path: Path,
    input_root: Path,
    output_dir: Path,
    sheet: str,
    *,
    policy: SecurityPolicy | None = None,
) -> Path:
    """Convert one explicitly selected XLSX sheet without deduplicating rows."""
    if not sheet:
        raise ExportFailure("SHEET_REQUIRED", "必须明确指定工作表")
    active_policy = policy or SecurityPolicy()
    source = active_policy.validate_input(Path(input_path), Path(input_root))
    if source.suffix.casefold() != ".xlsx":
        raise ExportFailure("SPREADSHEET_FORMAT", "表格操作目前只接受 XLSX")
    _reject_macros(source)

    service = SpreadsheetService(active_policy)
    summary = service.inspect(source, Path(input_root))
    if sheet not in _selected_sheet_names(summary):
        raise ExportFailure("SHEET_NOT_FOUND", "指定工作表不存在")

    # Scan the snapshot before creating any output.  A second read is used for
    # writing so a formula discovered late can never publish a partial CSV.
    with active_policy.snapshot_input(source, Path(input_root)) as staged:
        workbook = load_workbook(staged, read_only=True, data_only=False, keep_links=False)
        try:
            _scan_workbook(workbook, sheet)
        finally:
            workbook.close()

    output_root = active_policy.validate_output_dir(Path(output_dir), Path(input_root))
    target = safe_output_path(output_root, source.name, ".csv")
    temporary: Path | None = None
    try:
        with active_policy.snapshot_input(source, Path(input_root)) as staged:
            workbook = load_workbook(staged, read_only=True, data_only=False, keep_links=False)
            try:
                # Re-scan the source snapshot and write the exact worksheet rows
                # in source order, including duplicate and blank rows.
                rows = _scan_workbook(workbook, sheet)
                with tempfile.NamedTemporaryFile(
                    mode="wb", prefix=".financial-xlsx-", suffix=".tmp", dir=output_root, delete=False
                ) as handle:
                    temporary = Path(handle.name)
                _write_csv(rows, temporary)
            finally:
                workbook.close()
        publish_new_file(temporary, target)
        return target
    finally:
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Optional local XLSX-to-CSV preprocessor; platform runtime remains UNKNOWN."
    )
    parser.add_argument("--input", required=True, type=Path, help="one XLSX input file")
    parser.add_argument("--input-root", required=True, type=Path, help="allowed input root")
    parser.add_argument("--output-dir", required=True, type=Path, help="output directory outside input root")
    parser.add_argument("--sheet", required=True, help="exact worksheet name")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        output = convert_xlsx_to_csv(args.input, args.input_root, args.output_dir, args.sheet)
    except EngineError as exc:
        print(f"{exc.code}: {exc.public_message}", file=sys.stderr)
        return 1
    except Exception:
        print("XLSX_CONVERSION_ERROR: XLSX 转换失败，详细异常已脱敏", file=sys.stderr)
        return 1
    # Print only a basename: callers can already determine --output-dir, while
    # diagnostics must not echo local absolute paths or spreadsheet contents.
    print(output.name)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
