"""Safe, non-formula-executing XLSX inspection and deduplication."""
from __future__ import annotations

import csv
import time
import uuid
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .errors import EngineError, ExportFailure, SecurityViolation
from .metadata import publish_new_file, safe_output_path
from .models import ExportResult, SpreadsheetSummary
from .resource_slots import hold_resource_slot
from .security import SecurityPolicy


class SpreadsheetService:
    def __init__(self, policy: SecurityPolicy | None = None):
        self.policy = policy or SecurityPolicy()

    def inspect(self, path: Path, input_root: Path) -> SpreadsheetSummary:
        safe = self.policy.validate_input(Path(path), input_root)
        if safe.suffix.lower() != ".xlsx":
            raise ExportFailure("SPREADSHEET_FORMAT", "表格操作目前只接受 XLSX")
        sheets: list[dict[str, object]] = []
        total_rows = total_nulls = total_duplicates = 0
        with hold_resource_slot(), self.policy.snapshot_input(safe, input_root) as staged:
            workbook = load_workbook(staged, read_only=True, data_only=False, keep_links=False)
            budget = {"rows": 0, "cells": 0, "text": 0}
            deadline = time.monotonic() + self.policy.timeout_seconds
            try:
                for sheet in workbook.worksheets:
                    rows = self._bounded_rows(sheet, budget, deadline)
                    headers = [str(value) if value is not None else "" for value in next(rows, ())]
                    seen: set[tuple[object, ...]] = set()
                    row_count = null_count = duplicate_count = 0
                    for row in rows:
                        row_count += 1
                        identity = tuple(row)
                        if identity in seen:
                            duplicate_count += 1
                        else:
                            seen.add(identity)
                        null_count += sum(value is None or value == "" for value in row)
                    total_rows += row_count
                    total_nulls += null_count
                    total_duplicates += duplicate_count
                    sheets.append({"name": sheet.title, "headers": headers, "row_count": row_count,
                                   "null_cells": null_count, "duplicate_rows": duplicate_count})
            finally:
                workbook.close()
        return SpreadsheetSummary(safe.name, sheets, total_rows, total_nulls, total_duplicates)

    def deduplicate(
        self,
        path: Path,
        input_root: Path,
        output_dir: Path,
        keys: Sequence[str] | None = None,
        output_format: str = "xlsx",
    ) -> ExportResult:
        safe = self.policy.validate_input(Path(path), input_root)
        if safe.suffix.lower() != ".xlsx":
            return ExportResult("failed", None, safe.name, 0, 0, 0, {"code": "SPREADSHEET_FORMAT", "message": "表格操作目前只接受 XLSX"})
        try:
            output_root = self.policy.validate_output_dir(output_dir, input_root)
        except EngineError as exc:
            return ExportResult("failed", None, safe.name, 0, 0, 0, {"code": exc.code, "message": exc.public_message})
        extension = ".csv" if output_format.lower() == "csv" else ".xlsx"
        target = safe_output_path(output_root, safe.name, ".deduplicated" + extension)
        temporary = target.with_name(f".{target.stem}.{uuid.uuid4().hex}.tmp{extension}")
        try:
            with hold_resource_slot(), self.policy.snapshot_input(safe, input_root) as staged:
                workbook = load_workbook(staged, read_only=True, data_only=False, keep_links=False)
                try:
                    source_sheet = workbook.active
                    budget = {"rows": 0, "cells": 0, "text": 0}
                    deadline = time.monotonic() + self.policy.timeout_seconds
                    rows = self._bounded_rows(source_sheet, budget, deadline)
                    header = next(rows, None)
                    if header is None:
                        raise ExportFailure("EMPTY_SHEET", "工作表为空")
                    headers = [str(v) if v is not None else "" for v in header]
                    key_indexes = [headers.index(key) for key in keys] if keys else list(range(len(headers)))
                    seen: set[tuple[object, ...]] = set()
                    row_count = unique_count = 0
                    # The write-only writer must know the final dimension at its
                    # first append; a bounded first pass preserves max_row for
                    # read-only clients as well as validating before any output.
                    for row in rows:
                        row_count += 1
                        identity = tuple(row[index] if index < len(row) else None for index in key_indexes)
                        if identity not in seen:
                            seen.add(identity)
                            unique_count += 1
                    seen.clear()
                    budget = {"rows": 0, "cells": 0, "text": 0}
                    rows = self._bounded_rows(source_sheet, budget, deadline)
                    next(rows)
                    if extension == ".csv":
                        handle = temporary.open("w", newline="", encoding="utf-8-sig")
                        output = None
                        writer = csv.writer(handle)
                        writer.writerow([self._safe_value(value) for value in header])
                    else:
                        handle = None
                        output = Workbook(write_only=True)
                        sheet = output.create_sheet(source_sheet.title[:31] or "Sheet1")
                        dimension = f"A1:{get_column_letter(max(1, len(header)))}{unique_count + 1}"
                        sheet.calculate_dimension = lambda: dimension
                        sheet.append([self._safe_value(value) for value in header])
                    saved = False
                    try:
                        for row in rows:
                            identity = tuple(row[index] if index < len(row) else None for index in key_indexes)
                            if identity in seen:
                                continue
                            seen.add(identity)
                            values = [self._safe_value(value) for value in row]
                            if handle is not None:
                                writer.writerow(values)
                                if handle.tell() > self.policy.max_export_bytes:
                                    raise ExportFailure("EXPORT_TOO_LARGE", "导出文件超过大小限制")
                            else:
                                sheet.append(values)
                        if output is not None:
                            output.save(temporary)
                            saved = True
                    finally:
                        if handle is not None:
                            handle.close()
                        if output is not None:
                            if not saved:
                                try:
                                    sheet.close()
                                except Exception:
                                    pass
                                writer = getattr(sheet, "_writer", None)
                                if writer is not None and Path(writer.out).exists():
                                    writer.cleanup()
                            output.close()
                finally:
                    workbook.close()
                if temporary.stat().st_size > self.policy.max_export_bytes:
                    raise ExportFailure("EXPORT_TOO_LARGE", "导出文件超过大小限制")
                self._validate_export(temporary, extension)
                # Do not replace a target created since safe_output_path's check.
                publish_new_file(temporary, target)
                return ExportResult("success", str(target), safe.name, row_count, unique_count, row_count - unique_count)
        except EngineError as exc:
            return ExportResult("failed", None, safe.name, 0, 0, 0, {"code": exc.code, "message": exc.public_message})
        except (KeyError, ValueError) as exc:
            return ExportResult("failed", None, safe.name, 0, 0, 0, {"code": "INVALID_KEYS", "message": str(exc)})
        except Exception:
            return ExportResult("failed", None, safe.name, 0, 0, 0, {"code": "EXPORT_ERROR", "message": "表格导出失败，详细异常已脱敏"})
        finally:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass

    def _bounded_rows(self, sheet: object, budget: dict[str, int], deadline: float):
        if (sheet.max_row or 0) > self.policy.max_spreadsheet_rows:
            raise SecurityViolation("SPREADSHEET_ROW_LIMIT", "表格行数超过限制")
        if (sheet.max_column or 0) > self.policy.max_spreadsheet_cells:
            raise SecurityViolation("SPREADSHEET_CELL_LIMIT", "表格列数超过限制")
        for row in sheet.iter_rows(values_only=True):
            budget["rows"] += 1
            budget["cells"] += len(row)
            budget["text"] += sum(len(value.encode("utf-8")) for value in row if isinstance(value, str))
            if budget["rows"] > self.policy.max_spreadsheet_rows:
                raise SecurityViolation("SPREADSHEET_ROW_LIMIT", "表格行数超过限制")
            if budget["cells"] > self.policy.max_spreadsheet_cells or budget["text"] > self.policy.max_export_bytes:
                raise SecurityViolation("SPREADSHEET_CELL_LIMIT", "表格单元格或文本量超过限制")
            if time.monotonic() >= deadline:
                raise SecurityViolation("SPREADSHEET_TIMEOUT", "表格处理超时")
            yield row

    @staticmethod
    def _safe_value(value: object) -> object:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    @staticmethod
    def _validate_export(target: Path, extension: str) -> None:
        if not target.exists() or target.stat().st_size == 0:
            raise ExportFailure("INVALID_EXPORT", "导出文件为空")
        if extension == ".xlsx":
            check = load_workbook(target, read_only=True, data_only=False)
            check.close()
        else:
            with target.open("r", encoding="utf-8-sig", newline="") as handle:
                next(csv.reader(handle), None)
