"""Local deterministic input candidates, never project writes or factual verification.

Only exact catalog labels, explicit field IDs and the small reviewed alias map
are recognized. Ordinary prose/ambiguous units are not interpreted by an AI.
"""
from __future__ import annotations

from decimal import Decimal
from datetime import datetime, timezone
import csv
import hashlib
import io
import multiprocessing
import os
from pathlib import Path
import re
import time

from wangbin_document_engine.errors import EngineError
from wangbin_document_engine.security import SecurityPolicy

from .attachment_security import AttachmentError, formula_text, inspect_snapshot, stable_snapshot
from .errors import ContractError
from .provenance import make_record


_ALIASES = {
    "项目位置": "site.address", "经度": "site.longitude", "纬度": "site.latitude",
    "屋顶或土地总面积": "site.total_area",
    "屋顶总面积": "site.total_area", "土地总面积": "site.total_area",
    "可用面积": "site.usable_area", "直流装机": "system.dc_capacity",
    "交流容量": "system.ac_capacity", "并网接入": "grid.connection_point",
}
_CAPABILITIES = {"mode": "LOCAL_INPUT_CANDIDATES_ONLY", "ocr": False,
                 "network": False, "calculations": False, "public_untrusted_service": "BLOCKED"}
_LIMITATIONS = ["KNOWN_LIMIT：不是操作系统沙箱；Windows硬内存及后代进程隔离未解决。",
                "只提取确定性标签和显式字段表头；候选须人工确认，未核实原件真实性或授权。",
                "不提供OCR、自然语言推断、联网、公式计算或复杂版面重建。"]


def _base(path):
    return {"status": "FAILED", "attachment_id": None, "sha256": None,
            "source_name": path.name, "format": path.suffix.lower().lstrip("."),
            "candidates": [], "quality_notes": [], "locations": [], "error": None,
            "capabilities": dict(_CAPABILITIES), "limitations": list(_LIMITATIONS)}


class _Extractor:
    def __init__(self, catalog, attachment_id, policy):
        self.catalog, self.attachment_id, self.policy = catalog, attachment_id, policy
        self.candidates, self.notes, self.locations = [], [], []
        self.labels = {}
        for fid, field in catalog.fields.items():
            self.labels.setdefault(field["label_zh"], []).append(fid)
        self.seen_text = False
        self.currency = None
        self.currency_conflicted = False
        self.cells = self.rows = self.text_bytes = 0
        self.deadline = time.monotonic() + policy.timeout_seconds

    def budget(self, cells=0, rows=0, text=""):
        self.cells += cells
        self.rows += rows
        self.text_bytes += len(text.encode("utf-8"))
        if self.cells > min(self.policy.max_spreadsheet_cells, 20000):
            raise AttachmentError("CELL_LIMIT")
        if self.rows > self.policy.max_spreadsheet_rows:
            raise AttachmentError("ROW_LIMIT")
        if self.text_bytes > min(self.policy.max_export_bytes, 8 * 1024 * 1024):
            raise AttachmentError("TEXT_LIMIT")
        if time.monotonic() >= self.deadline:
            raise AttachmentError("PARSER_TIMEOUT")

    def envelope(self, fid, value, location):
        identity = hashlib.sha256((self.attachment_id + "|" + fid + "|" + location).encode()).hexdigest()[:24]
        return make_record(self.catalog, fid, value, source_kind="DOCUMENT_EXTRACTED",
                           field_instance_id="extract-" + identity, source_id=self.attachment_id,
                           source_url_or_attachment=self.attachment_id, source_section_or_cell=location,
                           location_status="EXACT", confidence_basis="确定性标签/表头匹配，待人工核验",
                           retrieved_at=datetime.now(timezone.utc).isoformat(),
                           missing_reason="未核实原始资料、适用期间、真实性及授权；须人工确认")

    def drain(self, rows):
        # Unsupported headers downgrade extraction, not security scanning.
        for row in rows:
            self.budget(cells=len(row), rows=1, text="".join(str(value or "") for value, _ in row))

    def table(self, rows, *, key_value=False):
        """Rows contain (raw value, actual cell location) tuples; no inferred columns."""
        rows = iter(rows)
        first = next(rows, None)
        if first is None:
            self.notes.append("EMPTY_TABLE")
            return
        if key_value and len(first) == 2 and self.header(first[0][0])[0] and not self.header(first[1][0])[0]:
            for row in _prepend(first, rows):
                self.budget(cells=len(row), rows=1, text="".join(str(v or "") for v, _ in row))
                if len(row) != 2:
                    self.notes.append("COMPLEX_TABLE")
                    continue
                fid, unit = self.header(row[0][0])
                if fid and not self.catalog.fields[fid].get("repeat_group"):
                    record = self.record(fid, row[1][0], row[1][1], unit=unit)
                    if record:
                        self.candidates.append(record)
                        self.locations.append(row[1][1])
                else:
                    self.notes.append("UNRECOGNIZED_OR_AMBIGUOUS_LABEL")
            self.seen_text = True
            return
        self.budget(cells=len(first), rows=1, text="".join(str(v or "") for v, _ in first))
        headers = [self.header(raw) for raw, _ in first]
        known = [fid for fid, _ in headers if fid]
        if len(known) != len(set(known)):
            self.notes.append("DUPLICATE_TABLE_HEADERS")
            self.drain(rows)
            return
        if any(fid is None for fid, _ in headers):
            self.notes.append("UNKNOWN_TABLE_COLUMNS")
        # Repeated columns must be full IDs; ambiguous Chinese child labels do
        # not acquire group context by guessing.
        groups = {self.catalog.fields[fid].get("repeat_group") for fid in known if self.catalog.fields[fid].get("repeat_group")}
        if groups and (len(groups) != 1 or any(not self.catalog.fields[fid].get("repeat_group") for fid in known)
                       or any(fid and str(raw).strip().split(" ")[0].split("(")[0].split("（")[0] != fid
                              for (raw, _), (fid, _) in zip(first, headers))):
            self.notes.append("MIXED_OR_NONEXPLICIT_GROUP_TABLE")
            self.drain(rows)
            return
        group = next(iter(groups), None)
        group_rows = []
        source_row_ids = set()
        first_location = None
        for row in rows:
            self.budget(cells=len(row), rows=1, text="".join(str(v or "") for v, _ in row))
            if not any(value is not None and str(value).strip() for value, _ in row):
                continue
            self.seen_text = True
            if len(row) != len(headers):
                self.notes.append("IRREGULAR_TABLE_ROW")
                continue
            location = row[0][1]
            first_location = first_location or location
            row_id = "row-" + hashlib.sha256((self.attachment_id + "|" + str(group) + "|" + location).encode()).hexdigest()[:20]
            fields = {}
            for (fid, unit), (raw, cell_location) in zip(headers, row):
                if not fid:
                    continue
                record = self.record(fid, raw, cell_location, unit=unit, row_id=row_id)
                if record:
                    if group:
                        fields[fid] = record
                    else:
                        self.candidates.append(record)
                    self.locations.append(cell_location)
            if group and fields:
                row_fid = group + ".row_id"
                # A supplied identifier is literal evidence, never repaired.
                if row_fid in fields:
                    row_id=fields[row_fid]["value"]
                    if re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,95}",row_id) is None:
                        self.notes.append("SOURCE_ROW_ID_INVALID_REQUIRES_REVIEW")
                    if row_id in source_row_ids:
                        self.notes.append("SOURCE_ROW_ID_DUPLICATE_REQUIRES_REVIEW")
                else:
                    fields[row_fid] = self.record(row_fid, row_id, location, row_id=row_id)
                    fields[row_fid]["confidence_basis"] = "原文未提供行ID；由附件摘要、字段组及真实源行定位派生的技术ID，非工程参数"
                source_row_ids.add(row_id)
                group_rows.append({"row_id": row_id, "fields": fields})
        if group_rows:
            self.candidates.append(self.envelope(group, group_rows, first_location))

    def header(self, label):
        label = str(label).strip()
        match = re.fullmatch(r"(.+?)\s*[（(]([^()（）]+)[）)]", label)
        key, unit = (match[1].strip(), match[2].strip()) if match else (label, None)
        if key in self.catalog.fields:
            return key, unit
        if key in _ALIASES:
            return _ALIASES[key], unit
        ids = self.labels.get(key, [])
        return (ids[0], unit) if len(ids) == 1 else (None, unit)

    def record(self, fid, raw, location, *, unit=None, row_id=None, location_status="EXACT"):
        field = self.catalog.fields[fid]
        kind = field["type"]
        if raw is None or str(raw).strip() == "":
            self.notes.append("EMPTY_FIELD")
            return None
        text = str(raw).strip()
        if len(text) > 4000:
            self.notes.append("FIELD_TEXT_LIMIT")
            return None
        try:
            if kind in {"number", "integer"}:
                match = re.fullmatch(r"([+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]{1,3})?)\s*(\S.*)?", text)
                if not match or len(match[1]) > 128:
                    raise ValueError
                value = Decimal(match[1])
                if not value.is_finite() or abs(value.adjusted()) > 300:
                    raise ValueError
                explicit = match[2].strip() if match[2] else None
                if explicit and unit and explicit != unit:
                    raise ValueError
                unit = explicit or unit
                if field["unit"] is not None and unit is None:
                    self.notes.append("EXPLICIT_UNIT_REQUIRED")
                    return None
                if kind == "integer":
                    if value != value.to_integral_value():
                        raise ValueError
                    value = int(value)
            elif kind == "boolean":
                if isinstance(raw, bool):
                    value = raw
                elif text in {"true", "false", "是", "否"}:
                    value = text in {"true", "是"}
                else:
                    raise ValueError
            elif kind == "string":
                value = str(raw) if fid.endswith((".row_id",".task_id",".period_id",".boundary_id")) else text
                allowed = field["constraints"].get("enum")
                if allowed is not None and value not in allowed:
                    raise ValueError
                if fid=="project.currency" and re.fullmatch(r"[A-Z]{3}",value):
                    if self.currency is not None and self.currency!=value:
                        self.currency_conflicted=True
                        self.notes.append("DOCUMENT_CURRENCY_CONFLICT")
                    self.currency=None if self.currency_conflicted else value
            else:
                self.notes.append("COMPLEX_VALUE_REQUIRES_TABLE")
                return None
            identity = hashlib.sha256((self.attachment_id + "|" + fid + "|" + location + "|" + str(row_id)).encode()).hexdigest()[:24]
            record = make_record(self.catalog, fid, value, original_unit=unit,currency=self.currency,
                                 source_kind="DOCUMENT_EXTRACTED", source_id=self.attachment_id,
                                 source_url_or_attachment=self.attachment_id, source_section_or_cell=location,
                                 field_instance_id="extract-" + identity, location_status=location_status,
                                 retrieved_at=datetime.now(timezone.utc).isoformat(),
                                 confidence_basis="确定性标签/表头匹配，待人工核验",
                                 missing_reason="未核实原始资料、适用期间、真实性及授权；须人工确认")
            return record
        except (ContractError, ValueError, ArithmeticError):
            self.notes.append("VALUE_OR_UNIT_UNRESOLVED")
            return None

    def line(self, text, location, *, location_status="EXACT"):
        if not text.strip():
            return
        self.seen_text = True
        match = re.fullmatch(r"\s*(?:#{1,6}\s+)?([^:：\n]+)[:：]\s*(.*?)\s*", text)
        if not match:
            return
        fid, unit = self.header(match[1])
        if fid is None:
            self.notes.append("UNRECOGNIZED_OR_AMBIGUOUS_LABEL")
        elif self.catalog.fields[fid].get("repeat_group"):
            self.notes.append("REPEATED_FIELD_REQUIRES_TABLE")
        else:
            record = self.record(fid, match[2], location, unit=unit, location_status=location_status)
            if record:
                self.candidates.append(record)
                self.locations.append(location)


def _prepend(first, rest):
    yield first
    yield from rest


def _xlsx_numeric_tokens(staged,policy):
    """Preserve OOXML decimal text before openpyxl casts it to binary float."""
    import posixpath
    import zipfile
    from defusedxml import ElementTree
    main="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    relationship="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    result={};count=0
    with zipfile.ZipFile(staged) as archive:
        book=ElementTree.fromstring(archive.read("xl/workbook.xml"),forbid_dtd=True)
        rels=ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"),forbid_dtd=True)
        targets={r.attrib["Id"]:r.attrib["Target"] for r in rels}
        for sheet in book.find(main+"sheets"):
            target=targets[sheet.attrib[relationship]]
            target=target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/"+target)
            if not target.startswith("xl/worksheets/"):
                raise AttachmentError("XLSX_INVALID_SHEET_REFERENCE")
            tree=ElementTree.fromstring(archive.read(target),forbid_dtd=True)
            numbers={}
            for cell in tree.iter(main+"c"):
                count+=1
                if count>min(policy.max_spreadsheet_cells,20000):
                    raise AttachmentError("CELL_LIMIT")
                value=cell.find(main+"v")
                if cell.attrib.get("t","n")=="n" and value is not None and value.text is not None:
                    numbers[cell.attrib["r"]]=value.text
            result[sheet.attrib["name"]]=numbers
    return result


def _extract(staged, catalog, attachment_id, policy, encoding):
    extractor = _Extractor(catalog, attachment_id, policy)
    suffix = staged.suffix
    needs_ocr = False
    if suffix in {".txt", ".md", ".csv"}:
        if encoding not in {"utf-8-sig", "utf-8", "gb18030", "utf-16", "utf-16-le", "utf-16-be"}:
            raise AttachmentError("UNSUPPORTED_ENCODING", "FAILED")
        text = staged.read_bytes().decode(encoding, errors="strict")
        if "\x00" in text or any(ord(c) < 32 and c not in "\t\r\n" for c in text):
            raise AttachmentError("BINARY_TEXT_INPUT", "FAILED")
        if suffix == ".csv":
            reader = csv.reader(io.StringIO(text, newline=""), strict=True)
            def csv_rows():
                previous = 0
                for row in reader:
                    if any(formula_text(value) for value in row):
                        raise AttachmentError("FORMULA_INPUT")
                    start, end = previous + 1, reader.line_num
                    previous = end
                    span = str(start) if start == end else f"{start}-{end}"
                    yield [(value, f"line:{span}/column:{col}") for col, value in enumerate(row, 1)]
            extractor.table(csv_rows())
        else:
            for number, line in enumerate(text.splitlines(), 1):
                extractor.budget(rows=1, cells=1, text=line)
                if "|" in line or "\t" in line:
                    extractor.notes.append("TEXT_TABLE_LAYOUT_NOT_RECONSTRUCTED")
                extractor.line(line, "line:" + str(number))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter, quote_sheetname
        numeric_tokens=_xlsx_numeric_tokens(staged,policy)
        workbook = load_workbook(staged, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > policy.max_spreadsheet_rows or (sheet.max_column or 0) > min(policy.max_spreadsheet_cells, 20000):
                    raise AttachmentError("SPREADSHEET_DIMENSION_LIMIT")
                def sheet_rows():
                    for row_number, row in enumerate(sheet.iter_rows(), 1):
                        if any(cell.data_type == "f" or formula_text(cell.value) for cell in row):
                            raise AttachmentError("FORMULA_INPUT")
                        yield [(numeric_tokens.get(sheet.title,{}).get(get_column_letter(column)+str(row_number),cell.value), f"{quote_sheetname(sheet.title)}!{get_column_letter(column)}{row_number}")
                               for column, cell in enumerate(row, 1)]
                extractor.table(sheet_rows())
        finally:
            workbook.close()
    elif suffix == ".docx":
        from docx import Document
        document = Document(staged)
        for number, paragraph in enumerate(document.paragraphs, 1):
            extractor.budget(rows=1, cells=1, text=paragraph.text)
            lines = paragraph.text.splitlines()
            for line_number, line in enumerate(lines, 1):
                location = f"paragraph:{number}" + (f"/line:{line_number}" if len(lines) > 1 else "")
                extractor.line(line, location)
        for number, table in enumerate(document.tables, 1):
            if table._tbl.xpath(".//w:gridSpan|.//w:vMerge|.//w:tbl//w:tbl"):
                extractor.notes.append("COMPLEX_DOCX_TABLE_NOT_RECONSTRUCTED")
                continue
            def docx_rows():
                for row_number, row in enumerate(table.rows, 1):
                    yield [(cell.text, f"table:{number}/row:{row_number}/cell:{column}")
                           for column, cell in enumerate(row.cells, 1)]
            extractor.table(docx_rows(), key_value=True)
        if document.inline_shapes:
            extractor.notes.append("IMAGE_CONTENT_REQUIRES_OCR")
        needs_ocr = not extractor.seen_text
    elif suffix == ".pdf":
        from pdfminer.high_level import extract_pages
        from pdfminer.layout import LTTextContainer
        for page_number, page in enumerate(extract_pages(staged), 1):
            extractor.budget(rows=1)
            for text_number, item in enumerate((item for item in page if isinstance(item, LTTextContainer)), 1):
                for line_number, line in enumerate(item.get_text().splitlines(), 1):
                    extractor.budget(cells=1, text=line)
                    extractor.line(line, f"page:{page_number}/text:{text_number}/line:{line_number}", location_status="TEXT_FRAGMENT")
        needs_ocr = not extractor.seen_text
        extractor.notes.append("PDF_LAYOUT_REQUIRES_REVIEW")
    notes = list(dict.fromkeys(extractor.notes))
    if not extractor.candidates:
        notes.append("NO_DETERMINISTIC_CANDIDATES")
    return {"status": "NEEDS_OCR" if needs_ocr else ("PARTIAL" if notes else "PARSED"),
            "candidates": extractor.candidates, "quality_notes": notes,
            "locations": list(dict.fromkeys(extractor.locations))}


def _worker(staged, fields, attachment_id, policy, encoding, sender):
    """One disposable parser process; no plugins, renderer or subprocess parser."""
    from contextlib import redirect_stderr, redirect_stdout
    import logging
    from types import SimpleNamespace
    from .serialization import dumps
    logging.disable(logging.CRITICAL)
    with open(os.devnull, "w") as sink, redirect_stderr(sink), redirect_stdout(sink):
        try:
            if os.name != "nt":
                import resource
                cpu = max(1, int(policy.max_cpu_seconds))
                resource.setrlimit(resource.RLIMIT_CPU, (cpu, cpu))
                resource.setrlimit(resource.RLIMIT_AS, (policy.max_memory_bytes, policy.max_memory_bytes))
            inspect_snapshot(Path(staged), policy)
            result = _extract(Path(staged), SimpleNamespace(fields=fields), attachment_id, policy, encoding)
        except (AttachmentError, EngineError) as exc:
            result = {"status": getattr(exc, "status", "BLOCKED"), "error": {"code": exc.code, "message": "附件被安全策略拒绝或无法受限解析；详细内容已脱敏。"}}
        except (UnicodeError, LookupError):
            result = {"status": "FAILED", "error": {"code": "INVALID_ENCODING", "message": "附件编码不符合明确指定的编码；未猜测或替换字符。"}}
        except ImportError:
            result = {"status": "FAILED", "error": {"code": "PARSER_DEPENDENCY_UNAVAILABLE", "message": "本地解析依赖不可用；没有自动安装依赖或联网。"}}
        except Exception:
            result = {"status": "FAILED", "error": {"code": "ATTACHMENT_PARSE_FAILED", "message": "附件解析失败；详细异常已脱敏。"}}
        try:
            payload = dumps(result).encode("utf-8")
            if len(payload) > min(policy.max_export_bytes, 8 * 1024 * 1024):
                raise ValueError
        except Exception:
            payload = dumps({"status": "BLOCKED", "error": {"code": "PARSER_OUTPUT_LIMIT", "message": "附件候选输出超过限制。"}}).encode("utf-8")
        try:
            sender.send_bytes(payload)
        finally:
            sender.close()


def _run_bounded(staged, catalog, attachment_id, policy, encoding):
    from .serialization import loads
    context = multiprocessing.get_context("spawn")
    receiver, sender = context.Pipe(duplex=False)
    process = context.Process(target=_worker, args=(str(staged), catalog.fields, attachment_id, policy, encoding, sender), daemon=True)
    deadline = time.monotonic() + policy.timeout_seconds
    try:
        process.start()
        sender.close()
        while True:
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                raise AttachmentError("PARSER_TIMEOUT")
            if receiver.poll(min(remaining, 0.05)):
                return loads(receiver.recv_bytes(maxlength=32 * 1024 * 1024))
            if not process.is_alive():
                raise AttachmentError("PARSER_PROCESS_FAILED", "FAILED")
    finally:
        receiver.close()
        sender.close()
        if process.pid is not None:
            process.join(timeout=0.1)
            if process.is_alive():
                process.terminate()
                process.join(timeout=2)
            if process.is_alive():
                process.kill()
                process.join(timeout=2)
            process.close()


def parse_attachment(path: Path, input_root: Path, *, catalog, work_root: Path | None = None,
                     policy: SecurityPolicy | None = None, encoding: str = "utf-8-sig") -> dict:
    """Return sanitized Chinese diagnostics and unverified evidence envelopes."""
    path = Path(path)
    result = _base(path)
    active = policy or SecurityPolicy()
    try:
        with stable_snapshot(path, input_root, work_root, active) as staged:
            payload = staged.read_bytes()
            sha = hashlib.sha256(payload).hexdigest()
            del payload
            result.update(sha256=sha, attachment_id="att-" + sha[:16])
            result.update(_run_bounded(staged, catalog, result["attachment_id"], active, encoding))
    except (AttachmentError, EngineError) as exc:
        result.update(status=getattr(exc, "status", "BLOCKED"),
                      error={"code": exc.code, "message": "附件被安全策略拒绝或无法受限解析；详细内容已脱敏。"})
    except (UnicodeError, LookupError):
        result.update(status="FAILED", error={"code": "INVALID_ENCODING", "message": "附件编码不符合明确指定的编码；未猜测或替换字符。"})
    except Exception:
        result.update(status="FAILED", error={"code": "ATTACHMENT_PARSE_FAILED", "message": "附件解析失败；详细异常已脱敏。"})
    return result
