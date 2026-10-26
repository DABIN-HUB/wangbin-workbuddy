"""Generate deterministic real office-document fixtures without mocks."""
from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def make_pdf(path: Path) -> None:
    stream = b"BT /F1 20 Tf 72 720 Td (WANGBIN PDF MARKER) Tj 0 -32 Td (Real PDF fixture) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects)+1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(output)


def make_fixtures() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    (FIXTURES / "sample.txt").write_text("WANGBIN TXT MARKER\nTask: review the document.\n", encoding="utf-8")
    (FIXTURES / "sample.md").write_text("# WANGBIN MARKDOWN MARKER\n\n## Key point\nThe source is real.\n", encoding="utf-8")
    (FIXTURES / "sample.html").write_text("<html><body><h1>WANGBIN HTML MARKER</h1><p>Real HTML fixture.</p></body></html>", encoding="utf-8")
    make_pdf(FIXTURES / "sample.pdf")

    document = Document()
    document.add_heading("WANGBIN DOCX MARKER", level=1)
    document.add_paragraph("Real DOCX fixture with a key action item.")
    document.add_table(rows=2, cols=2).rows[0].cells[0].text = "Header"
    document.tables[0].rows[0].cells[1].text = "Value"
    document.tables[0].rows[1].cells[0].text = "Action"
    document.tables[0].rows[1].cells[1].text = "Review"
    document.save(FIXTURES / "sample.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount", "Note"])
    sheet.append(["Alice", 1, "first"])
    sheet.append(["Alice", 1, "first"])
    sheet.append(["Bob", None, "=2+2"])
    workbook.save(FIXTURES / "sample.xlsx")
    workbook.close()

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "WANGBIN PPTX MARKER"
    slide.placeholders[1].text = "Real PPTX fixture with presentation text."
    presentation.save(FIXTURES / "sample.pptx")

    (FIXTURES / "empty.txt").write_bytes(b"")
    (FIXTURES / "damaged.pdf").write_bytes(b"%PDF-1.4\nthis is not a complete PDF")
    (FIXTURES / "mismatched.pdf").write_text("This is plain text with a PDF extension.", encoding="utf-8")
    (FIXTURES / "wrong.docx").write_text("This is not a DOCX container.", encoding="utf-8")
    (FIXTURES / "中文 文件.txt").write_text("WANGBIN CHINESE FILENAME MARKER", encoding="utf-8")
    (FIXTURES / "special [name] #1.txt").write_text("WANGBIN SPECIAL FILENAME MARKER", encoding="utf-8")
    (FIXTURES / "oversized.bin").write_bytes(b"x" * 1024)


if __name__ == "__main__":
    make_fixtures()
    print(f"generated real fixtures in {FIXTURES}")
