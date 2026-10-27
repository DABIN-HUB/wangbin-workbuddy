from __future__ import annotations

import csv
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from wangbin_document_engine import SecurityPolicy
from wangbin_document_engine.errors import EngineError
from tooling.financial_xlsx import convert_xlsx_to_csv, main


def make_workbook(path: Path, rows: list[list[object]], sheet: str = "Ledger") -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    workbook.create_sheet("Other").append(["not selected"])
    workbook.save(path)
    workbook.close()


def read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.reader(handle))


@pytest.mark.real_format
def test_real_xlsx_roundtrip_preserves_order_and_duplicate_rows(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    input_root.mkdir()
    source = input_root / "ledger.xlsx"
    make_workbook(
        source,
        [
            ["Entry", "Amount", "Currency"],
            ["A-1", 10, "CNY"],
            ["A-1", 10, "CNY"],
            ["B-2", 20.5, "USD"],
        ],
    )
    output_dir = tmp_path / "output"

    output = convert_xlsx_to_csv(source, input_root, output_dir, "Ledger")

    assert output.name == "ledger.csv"
    assert read_csv(output) == [
        ["Entry", "Amount", "Currency"],
        ["A-1", "10", "CNY"],
        ["A-1", "10", "CNY"],
        ["B-2", "20.5", "USD"],
    ]
    assert output.stat().st_size <= 1 * 1024 * 1024


def test_formula_cells_are_rejected_without_output(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    input_root.mkdir()
    source = input_root / "formula.xlsx"
    make_workbook(source, [["Amount"], ["=1+1"]])
    output_dir = tmp_path / "output"

    with pytest.raises(EngineError) as error:
        convert_xlsx_to_csv(source, input_root, output_dir, "Ledger")

    assert error.value.code == "FORMULA_INPUT"
    assert not output_dir.exists() or not list(output_dir.iterdir())


def test_macro_marked_xlsx_is_rejected(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    input_root.mkdir()
    source = input_root / "macro.xlsx"
    make_workbook(source, [["Amount"], [1]])
    temporary = tmp_path / "macro.tmp.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(temporary, "w") as rewritten:
        for item in original.infolist():
            rewritten.writestr(item, original.read(item.filename))
        rewritten.writestr("xl/vbaProject.bin", b"fictional macro payload")
    temporary.replace(source)

    with pytest.raises(EngineError) as error:
        convert_xlsx_to_csv(source, input_root, tmp_path / "output", "Ledger")

    assert error.value.code == "MACRO_INPUT"


def test_oversized_archive_is_rejected_before_output(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    input_root.mkdir()
    source = input_root / "large.xlsx"
    make_workbook(source, [["Value"], ["x" * 10_000]])

    with pytest.raises(EngineError) as error:
        convert_xlsx_to_csv(
            source,
            input_root,
            tmp_path / "output",
            "Ledger",
            policy=SecurityPolicy(max_uncompressed_bytes=256),
        )

    assert error.value.code in {"ZIP_ENTRY_TOO_LARGE", "ZIP_TOO_LARGE"}


def test_wrong_extension_and_empty_workbook_are_rejected(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    input_root.mkdir()
    wrong = input_root / "ledger.xls"
    wrong.write_bytes(b"not an xls workbook")
    with pytest.raises(EngineError) as wrong_error:
        convert_xlsx_to_csv(wrong, input_root, tmp_path / "output", "Ledger")
    assert wrong_error.value.code == "UNSUPPORTED_FORMAT"

    empty = input_root / "empty.xlsx"
    workbook = Workbook()
    workbook.save(empty)
    workbook.close()
    with pytest.raises(EngineError) as empty_error:
        convert_xlsx_to_csv(empty, input_root, tmp_path / "empty-output", "Sheet")
    assert empty_error.value.code == "EMPTY_SHEET"


def test_unsafe_input_path_is_rejected(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    input_root.mkdir()
    outside = tmp_path / "outside.xlsx"
    make_workbook(outside, [["Value"], [1]])

    with pytest.raises(EngineError) as error:
        convert_xlsx_to_csv(outside, input_root, tmp_path / "output", "Ledger")

    assert error.value.code == "PATH_OUTSIDE_ROOT"


def test_cli_privacy_does_not_print_absolute_paths_or_row_data(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    input_root = tmp_path / "private-client-directory"
    input_root.mkdir()
    source = input_root / "secret-client.xlsx"
    make_workbook(source, [["Amount"], ["=PRIVATE_ACCOUNT_NUMBER"]])

    status = main(
        [
            "--input",
            str(source),
            "--input-root",
            str(input_root),
            "--output-dir",
            str(tmp_path / "output"),
            "--sheet",
            "Ledger",
        ]
    )
    captured = capsys.readouterr()

    assert status == 1
    assert str(source) not in captured.out + captured.err
    assert "PRIVATE_ACCOUNT_NUMBER" not in captured.out + captured.err


def test_cli_requires_all_explicit_file_arguments() -> None:
    with pytest.raises(SystemExit):
        main([])
