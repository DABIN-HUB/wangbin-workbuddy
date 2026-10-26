from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from wangbin_document_engine import SecurityPolicy, SpreadsheetService

FIXTURES = Path(__file__).parent / "fixtures"


def test_spreadsheet_inspection_reports_duplicates_and_nulls() -> None:
    summary = SpreadsheetService().inspect(FIXTURES / "sample.xlsx", FIXTURES)
    assert summary.total_rows == 3
    assert summary.total_duplicate_rows == 1
    assert summary.total_null_cells == 1
    assert summary.worksheets[0]["headers"] == ["Name", "Amount", "Note"]


def test_deduplication_creates_reopenable_safe_xlsx(tmp_path: Path) -> None:
    result = SpreadsheetService().deduplicate(FIXTURES / "sample.xlsx", FIXTURES, tmp_path)
    assert result.status == "success", result.error
    assert result.duplicates_removed == 1
    output = Path(result.output_path)
    assert output.exists()
    workbook = load_workbook(output, data_only=False, read_only=True)
    assert workbook.active.max_row == 3
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    assert len(rows) == 3
    assert rows[-1][-1] == "'=2+2"


def test_csv_export_and_formula_safety(tmp_path: Path) -> None:
    result = SpreadsheetService().deduplicate(FIXTURES / "sample.xlsx", FIXTURES, tmp_path, output_format="csv")
    assert result.status == "success"
    assert Path(result.output_path).suffix == ".csv"
    assert "'=2+2" in Path(result.output_path).read_text(encoding="utf-8-sig")


def test_output_root_must_be_separate_from_input_root() -> None:
    result = SpreadsheetService().deduplicate(FIXTURES / "sample.xlsx", FIXTURES, FIXTURES)
    assert result.status == "failed"
    assert result.error["code"] == "OUTPUT_EQUALS_INPUT_ROOT"
