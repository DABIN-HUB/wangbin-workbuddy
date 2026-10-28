"""SYNTHETIC adversarial inputs, generated only beneath scratch pytest basetemp."""
from dataclasses import replace
import os
from pathlib import Path
import warnings
import zipfile

import pytest

from wangbin_document_engine.security import SecurityPolicy
from wangbin_solar_pv.catalog import Catalog
from wangbin_solar_pv.parsers import parse_attachment
from test_parsers import make_pdf


@pytest.fixture(scope="module")
def catalog():
    return Catalog()


def parse(path, catalog, *, root=None, **kwargs):
    return parse_attachment(path, root or path.parent, catalog=catalog, work_root=path.parent / "work", **kwargs)


def blocked(result, code=None):
    assert result["status"] == "BLOCKED", result
    assert result["candidates"] == []
    assert result["error"]["code"]
    if code:
        assert result["error"]["code"] == code
    assert "C:\\" not in str(result["error"])
    assert "SYNTHETIC-secret" not in str(result["error"])


def office(path, extra_name=None, extra_data=b"SYNTHETIC", *, document_xml=None):
    from docx import Document
    from openpyxl import Workbook
    if path.suffix == ".docx":
        doc = Document()
        doc.add_paragraph("项目名称：SYNTHETIC")
        doc.save(path)
    else:
        book = Workbook()
        book.active.append(["project.name"])
        book.active.append(["SYNTHETIC"])
        book.save(path)
        book.close()
    if document_xml is not None:
        with zipfile.ZipFile(path) as archive:
            entries = [(i.filename, archive.read(i)) for i in archive.infolist()]
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries:
                archive.writestr(name, document_xml if name == "word/document.xml" else data)
    if extra_name:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(extra_name, extra_data)


@pytest.mark.parametrize("suffix", [".txt", ".csv", ".docx", ".xlsx", ".pdf"])
def test_outside_root_is_blocked_for_every_format(tmp_path, catalog, suffix):
    root = tmp_path / "allowed"
    root.mkdir()
    path = tmp_path / ("SYNTHETIC-secret" + suffix)
    path.write_text("SYNTHETIC-secret", encoding="utf-8")
    blocked(parse(path, catalog, root=root), "PATH_OUTSIDE_ROOT")


def test_dot_dot_even_within_root_and_hardlink_are_blocked(tmp_path, catalog):
    sub = tmp_path / "sub"
    sub.mkdir()
    source = tmp_path / "SYNTHETIC.txt"
    source.write_text("project.name: SYNTHETIC", encoding="utf-8")
    blocked(parse(sub / ".." / source.name, catalog, root=tmp_path), "PATH_TRAVERSAL")
    linked = tmp_path / "hardlink.txt"
    os.link(source, linked)
    blocked(parse(linked, catalog), "HARDLINK_INPUT")


@pytest.mark.parametrize("kind", ["file", "root"])
def test_symlinks_and_linked_root_are_rejected_or_explicitly_skipped(tmp_path, catalog, kind):
    source = tmp_path / "SYNTHETIC.txt"
    source.write_text("project.name: SYNTHETIC", encoding="utf-8")
    link = tmp_path / "link"
    try:
        link.symlink_to(source if kind == "file" else tmp_path, target_is_directory=kind == "root")
    except (OSError, NotImplementedError) as exc:
        pytest.skip("Windows symlink privilege unavailable: " + type(exc).__name__)
    if kind == "file":
        blocked(parse(link, catalog))
    else:
        blocked(parse(link / source.name, catalog, root=link))


def test_oversize_encoding_and_timeout_are_explicit(tmp_path, catalog):
    path = tmp_path / "SYNTHETIC.txt"
    path.write_bytes(b"project.name: SYNTHETIC" * 20)
    blocked(parse(path, catalog, policy=replace(SecurityPolicy(), max_file_bytes=8)), "FILE_TOO_LARGE")
    blocked(parse(path, catalog, policy=replace(SecurityPolicy(), timeout_seconds=0)), "PARSER_TIMEOUT")
    path.write_bytes(b"\xff\xfe\x81broken")
    result = parse(path, catalog)
    assert result["status"] == "FAILED"
    assert result["error"]["code"] == "INVALID_ENCODING"
    path.write_bytes("项目名称：SYNTHETIC".encode("gb18030"))
    assert parse(path, catalog, encoding="gb18030")["status"] == "PARSED"
    assert parse(path, catalog, encoding="unicode_escape")["error"]["code"] == "UNSUPPORTED_ENCODING"


@pytest.mark.parametrize("name,payload,code", [
    ("word/vbaProject.bin", b"SYNTHETIC", "OFFICE_ACTIVE_CONTENT"),
    ("xl/macrosheets/sheet1.xml", b"<SYNTHETIC/>", "OFFICE_ACTIVE_CONTENT"),
    ("word/embeddings/oleObject1.bin", b"SYNTHETIC", "OFFICE_ACTIVE_CONTENT"),
    ("../SYNTHETIC.xml", b"SYNTHETIC", "ZIP_PATH_TRAVERSAL"),
    ("word/document.xml", b"SYNTHETIC", "ZIP_DUPLICATE_MEMBER"),
    ("WORD/DOCUMENT.XML", b"SYNTHETIC", "ZIP_DUPLICATE_MEMBER"),
    ("word/bomb.bin", b"0" * 4000000, "ZIP_COMPRESSION_RATIO"),
    ("word/_rels/evil.rels", b'<Relationships><Relationship TargetMode="External" Target="https://SYNTHETIC.invalid"/></Relationships>', "OFFICE_EXTERNAL_RELATIONSHIP"),
    ("word/active.xml", b'<x xmlns:w="urn:SYNTHETIC"><w:instrText>DDEAUTO SYNTHETIC</w:instrText></x>', "OFFICE_ACTIVE_CONTENT"),
], ids=["macro", "xlm", "ole", "traversal", "duplicate", "case-collision", "zipbomb", "external", "dde"])
def test_office_rejects_active_or_hostile_container(tmp_path, catalog, name, payload, code):
    path = tmp_path / "SYNTHETIC.docx"
    office(path, name, payload)
    blocked(parse(path, catalog), code)


def test_xml_entities_utf16_dde_and_container_limits(tmp_path, catalog):
    path = tmp_path / "SYNTHETIC.docx"
    office(path, document_xml=b'<!DOCTYPE x [<!ENTITY leak SYSTEM "file:///SYNTHETIC-secret">]><x>&leak;</x>')
    blocked(parse(path, catalog), "OFFICE_UNSAFE_XML")
    office(path, "word/dde.xml", '<x><instrText>DDEAUTO SYNTHETIC</instrText></x>'.encode("utf-16"))
    blocked(parse(path, catalog), "OFFICE_ACTIVE_CONTENT")
    office(path)
    blocked(parse(path, catalog, policy=replace(SecurityPolicy(), max_zip_entries=1)), "ZIP_TOO_MANY_ENTRIES")
    blocked(parse(path, catalog, policy=replace(SecurityPolicy(), max_uncompressed_bytes=100)), "ZIP_ENTRY_TOO_LARGE")


@pytest.mark.parametrize("value", ["=1+1", "  =1+1", "\t@SUM(A1:A2)", "+cmd|SYNTHETIC", " -2+3"])
def test_xlsx_formulas_and_whitespace_injection_blocked_in_unselected_sheet(tmp_path, catalog, value):
    from openpyxl import Workbook
    path = tmp_path / "SYNTHETIC.xlsx"
    book = Workbook()
    book.active.append(["project.name"])
    book.active.append(["SYNTHETIC"])
    book.create_sheet("hidden").append([value])
    book["hidden"].sheet_state = "hidden"
    book.save(path)
    book.close()
    blocked(parse(path, catalog), "FORMULA_INPUT")


def test_xlsx_real_numeric_negative_is_not_formula(tmp_path, catalog):
    from openpyxl import Workbook
    path = tmp_path / "SYNTHETIC.xlsx"
    book = Workbook()
    book.active.append(["site.latitude (degree)"])
    book.active.append([-25.2])
    book.save(path)
    book.close()
    result = parse(path, catalog)
    assert result["status"] == "PARSED", result
    assert result["candidates"][0]["value"] < 0


@pytest.mark.parametrize("extra", [b"/OpenAction << /S /JavaScript /JS (SYNTHETIC) >>",
                                     b"/Open#41ction << /S /Java#53cript /JS (SYNTHETIC) >>",
                                     b"/AA << /O << /S /Launch /F (SYNTHETIC) >> >>",
                                     b"/Names << /EmbeddedFiles << /Names [] >> >>"])
def test_pdf_structure_rejects_active_content_and_escaped_names(tmp_path, catalog, extra):
    path = tmp_path / "SYNTHETIC.pdf"
    make_pdf(path, catalog_extra=extra)
    blocked(parse(path, catalog), "PDF_ACTIVE_CONTENT")


def test_pdf_strings_are_not_scanned_as_executable_names(tmp_path, catalog):
    path = tmp_path / "SYNTHETIC.pdf"
    make_pdf(path, text="project.name: SYNTHETIC /JavaScript documentation")
    result = parse(path, catalog)
    assert result["status"] == "PARTIAL", result
    assert result["candidates"]


def test_csv_formula_injection_binary_and_row_budget(tmp_path, catalog):
    path = tmp_path / "SYNTHETIC.csv"
    path.write_text("project.name\n  =SYNTHETIC-secret\n", encoding="utf-8")
    blocked(parse(path, catalog), "FORMULA_INPUT")
    path.write_bytes(b"project.name\nSYNTHETIC\x00\n")
    assert parse(path, catalog)["error"]["code"] == "BINARY_TEXT_INPUT"
    path.write_text("project.name\nSYNTHETIC\nSYNTHETIC\n", encoding="utf-8")
    blocked(parse(path, catalog, policy=replace(SecurityPolicy(), max_spreadsheet_rows=2)), "ROW_LIMIT")


@pytest.mark.parametrize("duplicate_header", [False, True])
def test_xlsx_string_result_cell_cannot_hide_injection(tmp_path, catalog, duplicate_header):
    path = tmp_path / "SYNTHETIC.xlsx"
    office(path)
    with zipfile.ZipFile(path) as archive:
        entries = [(i.filename, archive.read(i)) for i in archive.infolist()]
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries:
            if name == "xl/worksheets/sheet1.xml":
                data = b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1:A2"/><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>project.name</t></is></c></row><row r="2"><c r="A2" t="str"><v>  =SYNTHETIC</v></c></row></sheetData></worksheet>'
                if duplicate_header:
                    data = data.replace(b'</row><row r="2">', b'<c r="B1" t="inlineStr"><is><t>project.name</t></is></c></row><row r="2">')
                    data = data.replace(b'ref="A1:A2"', b'ref="A1:B2"')
            archive.writestr(name, data)
    blocked(parse(path, catalog), "FORMULA_INPUT")
    csv = tmp_path / "SYNTHETIC.csv"
    csv.write_text("project.name,project.name\nSYNTHETIC, =SYNTHETIC\n", encoding="utf-8")
    blocked(parse(csv, catalog), "FORMULA_INPUT")


def test_private_snapshot_remains_identical_after_original_replacement(tmp_path, catalog):
    from wangbin_solar_pv.attachment_security import stable_snapshot, inspect_snapshot
    from wangbin_solar_pv.parsers import _extract
    path = tmp_path / "SYNTHETIC.txt"
    original = b"project.name: SYNTHETIC-original"
    path.write_bytes(original)
    with stable_snapshot(path, tmp_path, tmp_path / "work", SecurityPolicy()) as staged:
        # Windows retains the checked source handle; overwrite the content
        # rather than unlinking an in-use file (POSIX-only replacement idiom).
        path.write_text("project.name: SYNTHETIC-replaced", encoding="utf-8")
        assert staged.read_bytes() == original
        inspect_snapshot(staged, SecurityPolicy())
        result = _extract(staged, catalog, "att-SYNTHETIC", SecurityPolicy(), "utf-8-sig")
        assert result["candidates"][0]["value"] == "SYNTHETIC-original"
        staged_parent = staged.parent
    assert not staged_parent.exists()


def test_junction_root_rejected_with_real_windows_junction(tmp_path, catalog):
    if os.name != "nt":
        pytest.skip("Windows junction test")
    import subprocess
    target = tmp_path / "target"
    target.mkdir()
    path = target / "SYNTHETIC.txt"
    path.write_text("project.name: SYNTHETIC", encoding="utf-8")
    junction = tmp_path / "junction"
    result = subprocess.run(["cmd.exe", "/c", "mklink", "/J", str(junction), str(target)], capture_output=True)
    if result.returncode:
        pytest.skip("Windows junction creation unavailable")
    try:
        blocked(parse(junction / path.name, catalog, root=junction))
    finally:
        os.rmdir(junction)


def test_zip_encryption_flag_and_uncompressed_total_are_blocked(tmp_path, catalog):
    import struct
    path = tmp_path / "SYNTHETIC.docx"
    office(path)
    with zipfile.ZipFile(path) as archive:
        largest = max(i.file_size for i in archive.infolist())
        total = sum(i.file_size for i in archive.infolist())
    assert total > largest
    blocked(parse(path, catalog, policy=replace(SecurityPolicy(), max_uncompressed_bytes=largest)), "ZIP_TOO_LARGE")
    data = bytearray(path.read_bytes())
    for signature, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        position = data.index(signature)
        flags = struct.unpack_from("<H", data, position + offset)[0]
        struct.pack_into("<H", data, position + offset, flags | 1)
    path.write_bytes(data)
    blocked(parse(path, catalog), "ZIP_ENCRYPTED")


def test_spawn_timeout_kills_child_and_leaves_no_snapshot(tmp_path, catalog):
    import multiprocessing
    import time
    path = tmp_path / "SYNTHETIC.txt"
    path.write_text("project.name: SYNTHETIC", encoding="utf-8")
    initial = {p.pid for p in multiprocessing.active_children()}
    started = time.monotonic()
    blocked(parse(path, catalog, policy=replace(SecurityPolicy(), timeout_seconds=0.01)), "PARSER_TIMEOUT")
    assert time.monotonic() - started < 8
    assert {p.pid for p in multiprocessing.active_children()} == initial
    assert list((tmp_path / "work").iterdir()) == []


def test_compressed_pdf_object_stream_actions_are_blocked(tmp_path, catalog):
    """Real PDF 1.5 xref/object streams, compressed action absent from raw bytes."""
    import struct
    import zlib
    path = tmp_path / "SYNTHETIC.pdf"
    # Action object 6 is compressed inside object stream 7, indexed by xref 8.
    action = b"6 0 << /S /Java#53cript /JS (SYNTHETIC) >>"
    compressed = zlib.compress(action)
    text = b"BT /F1 12 Tf 30 720 Td (project.name: SYNTHETIC) Tj ET"
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R /SYNTHETIC 6 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        4: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        5: b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text + b"\nendstream",
        7: b"<< /Type /ObjStm /N 1 /First 4 /Filter /FlateDecode /Length " + str(len(compressed)).encode() + b" >>\nstream\n" + compressed + b"\nendstream",
    }
    data = bytearray(b"%PDF-1.5\n")
    offsets = {}
    for objid, payload in objects.items():
        offsets[objid] = len(data)
        data.extend(f"{objid} 0 obj\n".encode() + payload + b"\nendobj\n")
    offsets[8] = len(data)
    index = b"".join(struct.pack(">BIH", 0, 0, 65535) if i == 0 else struct.pack(">BIH", 2, 7, 0) if i == 6 else struct.pack(">BIH", 1, offsets[i], 0) for i in range(9))
    data.extend(b"8 0 obj\n<< /Type /XRef /Size 9 /Root 1 0 R /W [1 4 2] /Length " + str(len(index)).encode() + b" >>\nstream\n" + index + b"\nendstream\nendobj\n")
    data.extend(f"startxref\n{offsets[8]}\n%%EOF\n".encode())
    path.write_bytes(data)
    assert b"Java#53cript" not in data
    blocked(parse(path, catalog), "PDF_ACTIVE_CONTENT")
