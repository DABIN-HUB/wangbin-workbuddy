"""SYNTHETIC local documents only; every fixture is generated under pytest basetemp."""
from decimal import Decimal
import importlib.util
from pathlib import Path

import pytest

from wangbin_solar_pv.catalog import Catalog


@pytest.fixture(scope="module")
def catalog():
    return Catalog()


def parse(path, catalog, **kwargs):
    assert importlib.util.find_spec("wangbin_solar_pv.parsers") is not None, "Phase2D parser is not implemented"
    from wangbin_solar_pv.parsers import parse_attachment
    return parse_attachment(path, path.parent, catalog=catalog, work_root=path.parent / "work", **kwargs)


def test_text_labels_are_normalized_with_real_lines(tmp_path, catalog):
    path = tmp_path / "SYNTHETIC.txt"
    path.write_text("SYNTHETIC\n项目名称：SYNTHETIC屋顶\n直流装机：1.25 MWp\n贷款比例：70 %\n", encoding="utf-8")
    result = parse(path, catalog)
    assert result["status"] == "PARSED"
    records = {r["field_id"]: r for r in result["candidates"]}
    assert records["system.dc_capacity"]["value"] == Decimal("1250")
    assert records["finance.debt_ratio"]["value"] == Decimal("0.7")
    assert records["system.dc_capacity"]["source_section_or_cell"] == "line:3"
    assert records["system.dc_capacity"]["original_unit"] == "MWp"
    for record in records.values():
        assert record["status"] == "UNVERIFIED"
        assert record["license_status"] == "REQUIRES_REVIEW"
        assert record["source_id"] == result["attachment_id"]
        assert record["location_status"] == "EXACT"
    assert result["source_name"] == path.name
    assert result["attachment_id"] == "att-" + result["sha256"][:16]
    assert result["capabilities"]["public_untrusted_service"] == "BLOCKED"


def make_pdf(path, *, text="project.name: SYNTHETIC PDF", pages=1, catalog_extra=b""):
    """Real cross-referenced Helvetica PDF, no renderer or fake .pdf text."""
    stream = b"BT /F1 12 Tf 30 720 Td (" + text.encode("ascii") + b") Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R " + catalog_extra + b" >>",
               b"<< /Type /Pages /Kids [" + b" ".join(f"{3+i} 0 R".encode() for i in range(pages)) + b"] /Count " + str(pages).encode() + b" >>"]
    font, content = 3 + pages, 4 + pages
    objects.extend(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 " + str(font).encode() + b" 0 R >> >> /Contents " + str(content).encode() + b" 0 R >>" for _ in range(pages))
    objects.extend([b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
                    b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"])
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(output)


@pytest.mark.parametrize("suffix", [".md", ".csv", ".xlsx", ".docx", ".pdf"])
def test_real_formats_and_locations(tmp_path, catalog, suffix):
    path = tmp_path / ("SYNTHETIC" + suffix)
    expected = {".md": "line:2", ".csv": "line:2/column:1", ".xlsx": "'SYNTHETIC'!A2",
                ".docx": "paragraph:2", ".pdf": "page:1/text:1/line:1"}[suffix]
    if suffix == ".md":
        path.write_text("# SYNTHETIC\n项目名称：SYNTHETIC\n直流装机：2 MWp\n", encoding="utf-8")
    elif suffix == ".csv":
        path.write_text("project.name,system.dc_capacity (MWp)\nSYNTHETIC,2\n", encoding="utf-8")
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "SYNTHETIC"
        book.active.append(["project.name", "system.dc_capacity (MWp)"])
        book.active.append(["SYNTHETIC", 2])
        book.save(path)
        book.close()
    elif suffix == ".docx":
        from docx import Document
        document = Document()
        document.add_paragraph("SYNTHETIC")
        document.add_paragraph("项目名称：SYNTHETIC")
        table = document.add_table(rows=2, cols=2)
        for cell, value in zip(table.rows[0].cells, ["system.dc_capacity (MWp)", "finance.debt_ratio (%)"]):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ["2", "70"]):
            cell.text = value
        document.save(path)
    else:
        make_pdf(path, pages=2)
    result = parse(path, catalog)
    assert result["status"] in {"PARSED", "PARTIAL"}, result
    record = next(r for r in result["candidates"] if r["field_id"] == "project.name")
    assert record["source_section_or_cell"] == expected
    if suffix == ".pdf":
        assert record["location_status"] == "TEXT_FRAGMENT"
        assert any(r["source_section_or_cell"].startswith("page:2/") for r in result["candidates"])
    else:
        dc = next(r for r in result["candidates"] if r["field_id"] == "system.dc_capacity")
        assert dc["value"] == Decimal(2000)
        if suffix == ".docx":
            assert dc["source_section_or_cell"] == "table:1/row:2/cell:1"


@pytest.mark.parametrize("group,columns,values", [
    ("load.series", ["timestamp", "interval_minutes (minute)", "energy (kWh)"], ["2026-01-01T00:00:00+08:00", "60", "20"]),
    ("load.monthly_bills", ["period_start", "period_end", "energy (MWh)"], ["2026-01-01", "2026-01-31", "2"]),
    ("capex.items", ["category", "description", "amount (currency)", "tax_included"], ["modules", "SYNTHETIC quote", "10000", "true"]),
    ("finance.debt_schedule", ["date", "principal (currency)", "interest (currency)"], ["2026-01-01", "20", "1"]),
])
@pytest.mark.parametrize("suffix", [".csv", ".xlsx", ".docx"])
def test_container_rows_have_stable_ids_and_atomic_cell_evidence(tmp_path, catalog, group, columns, values, suffix):
    path = tmp_path / ("SYNTHETIC" + suffix)
    headers = [group + "." + col for col in columns]
    if suffix == ".csv":
        path.write_text(",".join(headers) + "\n" + ",".join(values) + "\n" + ",".join(values) + "\n", encoding="utf-8")
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "SYNTHETIC"
        for row in [headers, values, values]:
            book.active.append(row)
        book.save(path)
        book.close()
    else:
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=3, cols=len(headers))
        for row, data in zip(table.rows, [headers, values, values]):
            for cell, value in zip(row.cells, data):
                cell.text = value
        doc.save(path)
    result = parse(path, catalog)
    assert result["status"] == "PARSED", result
    parent = next(r for r in result["candidates"] if r["field_id"] == group)
    rows = parent["value"]
    assert len(rows) == 2
    assert rows[0]["row_id"] != rows[1]["row_id"]
    ids = {parent["field_instance_id"]}
    for row in rows:
        assert row["fields"][group + ".row_id"]["value"] == row["row_id"]
        for record in row["fields"].values():
            assert record["source_url_or_attachment"] == result["attachment_id"]
            location = record["source_section_or_cell"]
            assert location.startswith("line:") if suffix == ".csv" else ("!" in location if suffix == ".xlsx" else location.startswith("table:1/row:"))
            assert record["status"] == "UNVERIFIED"
            assert record["field_instance_id"] not in ids
            ids.add(record["field_instance_id"])
    again = parse(path, catalog)
    repeated = next(r for r in again["candidates"] if r["field_id"] == group)
    assert [r["row_id"] for r in rows] == [r["row_id"] for r in repeated["value"]]
    assert [r["field_instance_id"] for r in rows[0]["fields"].values()] == [r["field_instance_id"] for r in repeated["value"][0]["fields"].values()]


def test_ten_groups_and_strict_booleans_enums(tmp_path, catalog):
    path = tmp_path / "SYNTHETIC.txt"
    path.write_text("\n".join([
        "项目名称：SYNTHETIC", "项目位置：SYNTHETIC地址", "可用面积：2 公顷",
        "直流装机：2 MWp", "交流容量：1.5 MWac", "全年用电量：10 MWh",
        "组件型号：SYNTHETIC-module", "估算依据与价格基年：SYNTHETIC报价2026",
        "并网接入：SYNTHETIC接入点", "贷款比例：70 %", "是否配储：否",
        "project.pv_type: industrial_rooftop", "project.has_storage: maybe",
        "project.pv_type: 随便一种", "容量：20 MW", "system.dc_capacity: 10 MW",
        "site.total_area: 500", "纬度：30 degree"]), encoding="utf-8")
    result = parse(path, catalog)
    assert result["status"] == "PARTIAL"
    records = {r["field_id"]: r for r in result["candidates"]}
    assert set(["project.name", "site.address", "site.usable_area", "system.dc_capacity", "system.ac_capacity",
                "load.annual_energy", "system.module_model", "capex.estimate_basis", "grid.connection_point", "finance.debt_ratio"]) <= set(records)
    assert records["site.usable_area"]["value"] == 20000
    assert records["system.dc_capacity"]["value"] == 2000
    assert records["project.has_storage"]["value"] is False
    assert records["project.pv_type"]["value"] == "industrial_rooftop"
    assert "site.total_area" not in records


def test_unknown_columns_are_partial_not_silently_absent(tmp_path, catalog):
    path = tmp_path / "SYNTHETIC.csv"
    path.write_text('project.name,UNKNOWN_COLUMN\nSYNTHETIC,private-unused\n', encoding="utf-8")
    result = parse(path, catalog)
    assert result["status"] == "PARTIAL"
    assert "UNKNOWN_TABLE_COLUMNS" in result["quality_notes"]
    assert "private-unused" not in str(result)


def test_blank_pdf_requires_ocr_and_free_prose_does_not_make_up_values(tmp_path, catalog):
    pdf = tmp_path / "SYNTHETIC.pdf"
    make_pdf(pdf, text="")
    result = parse(pdf, catalog)
    assert result["status"] == "NEEDS_OCR"
    assert result["candidates"] == []
    path = tmp_path / "SYNTHETIC.txt"
    path.write_text("SYNTHETIC 屋顶差不多很大，装机可能一兆瓦，电价以后再说。", encoding="utf-8")
    result = parse(path, catalog)
    assert result["status"] == "PARTIAL"
    assert result["candidates"] == []


def test_docx_key_value_table_and_multiline_csv_real_lines(tmp_path, catalog):
    from docx import Document
    path = tmp_path / "SYNTHETIC.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "项目名称", "SYNTHETIC"
    table.cell(1, 0).text, table.cell(1, 1).text = "直流装机", "2 MWp"
    doc.save(path)
    result = parse(path, catalog)
    assert result["status"] == "PARSED", result
    assert {r["source_section_or_cell"] for r in result["candidates"]} == {"table:1/row:1/cell:2", "table:1/row:2/cell:2"}
    csv = tmp_path / "SYNTHETIC.csv"
    csv.write_text('project.name\n"SYNTHETIC\nwrapped"\nSYNTHETIC-next\n', encoding="utf-8")
    result = parse(csv, catalog)
    assert [r["source_section_or_cell"] for r in result["candidates"]] == ["line:2-3/column:1", "line:4/column:1"]


def test_xlsx_native_boolean_is_accepted_without_truthiness_coercion(tmp_path, catalog):
    from openpyxl import Workbook
    path = tmp_path / "SYNTHETIC.xlsx"
    book = Workbook()
    book.active.append(["project.has_storage"])
    book.active.append([True])
    book.active.append([False])
    book.active.append(["yes-ish"])
    book.save(path)
    book.close()
    result = parse(path, catalog)
    assert result["status"] == "PARTIAL"
    assert [r["value"] for r in result["candidates"]] == [True, False]


def test_docx_multiline_paragraph_instances_are_unique_and_coordinates_are_explicit(tmp_path, catalog):
    from docx import Document
    path = tmp_path / "SYNTHETIC.docx"
    doc = Document()
    doc.add_paragraph("项目名称：SYNTHETIC-one\n项目名称：SYNTHETIC-two\n经度：120 degree\n纬度：30 degree")
    doc.save(path)
    result = parse(path, catalog)
    assert len(result["candidates"]) == 4, result
    ids = [r["field_instance_id"] for r in result["candidates"]]
    assert len(ids) == len(set(ids))
    assert {r["field_id"] for r in result["candidates"]} >= {"site.latitude", "site.longitude"}


def test_complex_docx_merge_and_markdown_table_are_partial(tmp_path, catalog):
    from docx import Document
    path = tmp_path / "SYNTHETIC.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "SYNTHETIC merged"
    table.cell(1, 0).text = "项目名称"
    table.cell(1, 1).text = "SYNTHETIC"
    doc.save(path)
    result = parse(path, catalog)
    assert result["status"] in {"PARTIAL", "NEEDS_OCR"}
    assert "COMPLEX_DOCX_TABLE_NOT_RECONSTRUCTED" in result["quality_notes"]
    md = tmp_path / "SYNTHETIC.md"
    md.write_text("项目名称：SYNTHETIC\n| 时间 | 电量 |\n| 2026 | 10 |\n", encoding="utf-8")
    result = parse(md, catalog)
    assert result["status"] == "PARTIAL"
    assert "TEXT_TABLE_LAYOUT_NOT_RECONSTRUCTED" in result["quality_notes"]
